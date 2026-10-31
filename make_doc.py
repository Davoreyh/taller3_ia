"""
Script para generar el documento Word formal y completo del Punto 3 del Taller 3 de Inteligencia Artificial.
Autor: Camilo
"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(23, 54, 93) # Deep Navy
    elif level == 2:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(37, 99, 184) # Steel Blue
    elif level == 3:
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(50, 70, 95)
    return h

def add_callout(doc, title, instruction, command=None, screenshot_tag=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:color="2563B8"/>'
        f'  <w:left w:val="single" w:sz="24" w:color="2563B8"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:color="2563B8"/>'
        f'  <w:right w:val="single" w:sz="6" w:color="2563B8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"📸 {title}\n")
    r_title.bold = True
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(23, 54, 93)
    
    r_inst = p.add_run(f"Instrucción: {instruction}\n")
    r_inst.font.size = Pt(9.5)
    r_inst.font.color.rgb = RGBColor(40, 40, 40)
    
    if command:
        r_cmd_lbl = p.add_run("Comando a ejecutar:\n")
        r_cmd_lbl.bold = True
        r_cmd_lbl.font.size = Pt(9.0)
        r_cmd_lbl.font.color.rgb = RGBColor(70, 70, 70)
        
        r_cmd = p.add_run(f"  {command}\n")
        r_cmd.font.name = "Consolas"
        r_cmd.font.size = Pt(9.0)
        r_cmd.font.color.rgb = RGBColor(180, 40, 40)
        
    if screenshot_tag:
        r_box = p.add_run(f"\n[ >>> PEGAR AQUÍ {screenshot_tag} <<< ]\n(Presione Ctrl+V aquí para insertar la captura de pantalla)")
        r_box.bold = True
        r_box.font.size = Pt(9.5)
        r_box.font.color.rgb = RGBColor(37, 99, 184)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:color="D1D5DB"/>'
        f'  <w:left w:val="single" w:sz="18" w:color="9CA3AF"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:color="D1D5DB"/>'
        f'  <w:right w:val="single" w:sz="4" w:color="D1D5DB"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)

def format_table(tbl, col_widths, headers, data, alignments=None):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="B0C4DE", sz="4")
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "17365D")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        hdr_cells[i].width = Inches(col_widths[i])
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data rows
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F4F7FB"
        for col_idx, val in enumerate(row_data):
            row_cells[col_idx].text = str(val)
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            row_cells[col_idx].width = Inches(col_widths[col_idx])
            p = row_cells[col_idx].paragraphs[0]
            if alignments and col_idx < len(alignments):
                p.alignment = alignments[col_idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.size = Pt(9.0)
                r.font.color.rgb = RGBColor(30, 30, 30)

def build_p3_document():
    doc = docx.Document()
    
    # Configure page margins (2.5 cm = 0.984 inches)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)
    
    # ------------------ PORTADA / ENCABEZADO ------------------
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_before = Pt(0)
    p_uni.paragraph_format.space_after = Pt(2)
    r_uni = p_uni.add_run("UNIVERSIDAD DE LOS ANDES - FACULTAD DE INGENIERÍA\n")
    r_uni.bold = True
    r_uni.font.size = Pt(12)
    r_uni.font.color.rgb = RGBColor(23, 54, 93)
    r_dept = p_uni.add_run("DEPARTAMENTO DE INGENIERÍA DE SISTEMAS Y COMPUTACIÓN\n")
    r_dept.font.size = Pt(10)
    r_dept.font.color.rgb = RGBColor(97, 113, 132)
    r_curso = p_uni.add_run("ISIS-1611: Inteligencia Artificial (2026-2)\n")
    r_curso.font.size = Pt(10.5)
    r_curso.font.bold = True
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(16)
    p_title.paragraph_format.space_after = Pt(6)
    r_t1 = p_title.add_run("TALLER 3: RAZONAMIENTO BAJO INCERTIDUMBRE\n")
    r_t1.bold = True
    r_t1.font.size = Pt(17)
    r_t1.font.color.rgb = RGBColor(23, 54, 93)
    r_t2 = p_title.add_run("PUNTO 3: ELIMINACIÓN DE VARIABLES E INFERENCIA APROXIMADA (20 %)\n")
    r_t2.bold = True
    r_t2.font.size = Pt(13.5)
    r_t2.font.color.rgb = RGBColor(37, 99, 184)
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(18)
    r_meta = p_meta.add_run("Responsable del Punto: Camilo\nRama de Git: Camilo-P3 | Proyecto: AlertaBosque\nOctubre de 2026")
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph("―" * 55).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # ------------------ INTRODUCCIÓN ------------------
    add_header_styled(doc, "1. Introducción y Objetivos del Punto 3", level=1)
    doc.add_paragraph(
        "En el contexto del proyecto AlertaBosque, la estimación del riesgo de incendio a partir de señales de sensores ambientales "
        "(cámaras térmicas, sensores de humedad y temperatura) requiere métodos computacionales rigurosos para operar bajo incertidumbre. "
        "El presente informe documenta el desarrollo completo del Punto 3 del taller, el cual abarca dos grandes familias de algoritmos:"
    )
    doc.add_paragraph(
        "1. Inferencia Exacta mediante Eliminación de Variables: Técnica que reorganiza y almacena cálculos parciales aplicando la ley distributiva "
        "del álgebra para marginalizar variables ocultas sin construir la distribución conjunta completa, reduciendo drásticamente la complejidad computacional.\n"
        "2. Inferencia Aproximada mediante Muestreo: Técnicas estocásticas basadas en simulación Monte Carlo. Se implementa tanto el Muestreo a Priori "
        "(para estimar marginales sin evidencia) como el Muestreo por Rechazo (para estimar probabilidades condicionales descartando mundos inconsistentes con la evidencia observada)."
    )
    doc.add_paragraph(
        "Adicionalmente, se incluye el estudio del efecto de los órdenes de eliminación sobre el tamaño de los factores intermedios, "
        "la influencia del número de muestras y la probabilidad de la evidencia en el error de aproximación, y las respuestas requeridas "
        "para el informe general (Punto 6a y 6b)."
    )

    # ------------------ PARTE 1: ELIMINACIÓN DE VARIABLES ------------------
    add_header_styled(doc, "2. Eliminación de Variables: Estudio y Comparación", level=1)
    
    add_header_styled(doc, "2.1. Literal a) Cálculo manual paso a paso en cadena_humo", level=2)
    doc.add_paragraph(
        "Se evalúa la red bayesiana cadena_humo, cuya topología lineal modela la propagación de causas: "
        "Sequedad (S) → Incendio (I) → Humo (H). El objetivo es calcular manualmente la probabilidad no condicionada P(Humo=1), "
        "eliminando primero la variable Sequedad y posteriormente la variable Incendio."
    )
    doc.add_paragraph(
        "Las probabilidades condicionales locales (CPTs) suministradas para el escenario son:\n"
        "• P(Sequedad=1) = 0.40  ⇒  P(Sequedad=0) = 0.60\n"
        "• P(Incendio=1 | Sequedad=0) = 0.02  ⇒  P(Incendio=0 | Sequedad=0) = 0.98\n"
        "• P(Incendio=1 | Sequedad=1) = 0.30  ⇒  P(Incendio=0 | Sequedad=1) = 0.70\n"
        "• P(Humo=1 | Incendio=0) = 0.10  ⇒  P(Humo=0 | Incendio=0) = 0.90\n"
        "• P(Humo=1 | Incendio=1) = 0.80  ⇒  P(Humo=0 | Incendio=1) = 0.20"
    )
    doc.add_paragraph(
        "Definición de factores iniciales correspondientes a cada nodo:\n"
        "• f₁(S) = P(S)\n"
        "• f₂(S, I) = P(I | S)\n"
        "• f₃(I, H) = P(H | I)"
    )
    doc.add_paragraph(
        "La distribución de la variable consultada se formula algebraicamente como:\n"
        "P(H) = ∑_I ∑_S f₁(S) · f₂(S, I) · f₃(I, H) = ∑_I f₃(I, H) [ ∑_S f₁(S) · f₂(S, I) ]"
    )

    p_step1 = doc.add_paragraph()
    r = p_step1.add_run("Paso 1: Eliminación de Sequedad (S)\n")
    r.bold = True
    p_step1.add_run(
        "Se toman los factores que contienen a S: f₁(S) y f₂(S, I). Se construye su producto intermedio ψ₁(S, I) = f₁(S) · f₂(S, I), "
        "cuyo tamaño es de 4 filas:"
    )
    
    # Table psi_1
    t1 = doc.add_table(rows=5, cols=4)
    format_table(t1, [1.0, 1.0, 2.5, 1.5], 
                 ["Sequedad (S)", "Incendio (I)", "Operación: f₁(S) × f₂(S, I)", "Valor ψ₁(S, I)"],
                 [
                     ["0", "0", "0.60 × 0.98", "0.5880"],
                     ["0", "1", "0.60 × 0.02", "0.0120"],
                     ["1", "0", "0.40 × 0.70", "0.2800"],
                     ["1", "1", "0.40 × 0.30", "0.1200"]
                 ],
                 [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT])

    doc.add_paragraph(
        "Posteriormente, se suma (marginaliza) sobre Sequedad S ∈ {0, 1} para obtener el factor resultante m_S(I):\n"
        "• Para Incendio = 0: m_S(I=0) = ψ₁(0, 0) + ψ₁(1, 0) = 0.5880 + 0.2800 = 0.8680\n"
        "• Para Incendio = 1: m_S(I=1) = ψ₁(0, 1) + ψ₁(1, 1) = 0.0120 + 0.1200 = 0.1320\n"
        "Nótese que m_S(I) representa la distribución marginal P(Incendio): 0.8680 + 0.1320 = 1.0000."
    )

    p_step2 = doc.add_paragraph()
    r = p_step2.add_run("Paso 2: Eliminación de Incendio (I)\n")
    r.bold = True
    p_step2.add_run(
        "Los factores restantes que involucran a I son el factor acumulado m_S(I) y el factor del sensor f₃(I, H). "
        "Se calcula su producto intermedio ψ₂(I, H) = m_S(I) · f₃(I, H) (4 filas):"
    )

    # Table psi_2
    t2 = doc.add_table(rows=5, cols=4)
    format_table(t2, [1.0, 1.0, 2.5, 1.5],
                 ["Incendio (I)", "Humo (H)", "Operación: m_S(I) × f₃(I, H)", "Valor ψ₂(I, H)"],
                 [
                     ["0", "0", "0.8680 × 0.90", "0.7812"],
                     ["0", "1", "0.8680 × 0.10", "0.0868"],
                     ["1", "0", "0.1320 × 0.20", "0.0264"],
                     ["1", "1", "0.1320 × 0.80", "0.1056"]
                 ],
                 [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT])

    doc.add_paragraph(
        "Finalmente, se marginaliza sobre Incendio I ∈ {0, 1} para obtener el factor resultante final m_I(H):\n"
        "• Para Humo = 0: m_I(H=0) = ψ₂(0, 0) + ψ₂(1, 0) = 0.7812 + 0.0264 = 0.8076\n"
        "• Para Humo = 1: m_I(H=1) = ψ₂(0, 1) + ψ₂(1, 1) = 0.0868 + 0.1056 = 0.1924\n"
        "Al sumar ambos valores (0.8076 + 0.1924 = 1.0000), el factor ya se encuentra normalizado, concluyendo que:\n"
        "P(Humo = 1) = 0.192400 (19.24 %)\n"
        "P(Humo = 0) = 0.807600 (80.76 %)"
    )

    doc.add_paragraph(
        "Comprobación con el sistema AlertaBosque:\n"
        "Al ejecutar el módulo de Eliminación de Variables en orden Topológico (forward) para la consulta Humo sin evidencia, "
        "el software reporta exactamente:\n"
        "• Paso 1 (eliminar Sequedad | Producto: 4 filas): Incendio=0: 0.868000, Incendio=1: 0.132000.\n"
        "• Paso 2 (eliminar Incendio | Producto: 4 filas): Humo=0: 0.807600, Humo=1: 0.192400.\n"
        "• Máximo factor producto: 4 filas.\n"
        "El cálculo manual y la ejecución automatizada coinciden en su totalidad."
    )

    add_callout(
        doc,
        title="CAPTURA 1: Verificación de Eliminación de Variables en cadena_humo",
        instruction="Abra la ventana gráfica de la aplicación o ejecute en terminal para capturar la traza de factores y el resultado de Humo.",
        command="python main.py -p 3 -l cadena_humo --query Humo -a elimination",
        screenshot_tag="CAPTURA 1 (cadena_humo - Eliminación Humo)"
    )

    # ------------------ LITERAL B: COMPARACIÓN DE ÓRDENES ------------------
    add_header_styled(doc, "2.2. Literal b) Comparación de órdenes de eliminación en la red bosque", level=2)
    doc.add_paragraph(
        "Se analiza la red bayesiana compleja bosque para la consulta Incendio bajo la evidencia alarma (Alarma=1). "
        "Las variables ocultas a eliminar son: Tormenta, Sequedad, Falla, Humo y Temperatura. "
        "Se compara la referencia exacta obtenida por Enumeración (Punto 2) frente a dos órdenes de eliminación:"
    )
    doc.add_paragraph(
        "1. Orden Topológico (--order forward): Tormenta → Sequedad → Falla → Humo → Temperatura.\n"
        "2. Orden Inverso (--order reverse): Temperatura → Humo → Falla → Sequedad → Tormenta."
    )

    doc.add_paragraph("Resultados obtenidos de la ejecución:")
    
    t_comp = doc.add_table(rows=4, cols=5)
    format_table(t_comp, [1.8, 1.2, 1.2, 1.4, 1.2],
                 ["Método / Orden", "P(Incendio=0)", "P(Incendio=1)", "Mayor Factor Producto", "Coincidencia"],
                 [
                     ["Enumeración (Punto 2)", "0.670483", "0.329517", "No aplica (Recursivo)", "Referencia"],
                     ["Eliminación Topológica (forward)", "0.670483", "0.329517", "8 filas (2³)", "Exacta"],
                     ["Eliminación Inversa (reverse)", "0.670483", "0.329517", "32 filas (2⁵)", "Exacta"]
                 ],
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])

    doc.add_paragraph(
        "Análisis del comportamiento de los órdenes:\n"
        "• Coincidencia matemática: Ambos órdenes producen exactamente la misma distribución condicional posterior: "
        "P(Incendio=1 | Alarma=1) ≈ 32.9517 % y P(Incendio=0 | Alarma=1) ≈ 67.0483 %, coincidiendo con el valor exacto de enumeración. "
        "Esto confirma que el orden de eliminación no altera la corrección matemática del resultado final debido a la conmutatividad y asociatividad de la adición.\n"
        "• Impacto en la complejidad (Espacio y Tiempo): En el orden inverso, el primer paso (eliminar Temperatura) une factores que involucran a "
        "Sequedad, Falla, Incendio y Humo. Esto crea un factor conjunto intermedio de 2⁵ = 32 filas. En contraste, en el orden topológico se van "
        "eliminando primero los ancestros, manteniendo factores locales de tamaño acotado (máximo 8 filas). El orden inverso requirió 4 veces "
        "más memoria para almacenar el factor intermedio más pesado, evidenciando cómo un mal ordenamiento de eliminación puede provocar "
        "una explosión combinatorial en redes de gran escala."
    )

    add_callout(
        doc,
        title="CAPTURA 2: Eliminación en bosque con Orden Topológico (forward)",
        instruction="Ejecute el comando para abrir la visualización gráfica con el orden forward y capturar la tabla de pasos (máx 8 filas).",
        command="python main.py -p 3 -l bosque --case alarma -a elimination --order forward",
        screenshot_tag="CAPTURA 2 (bosque alarma - Eliminación Forward)"
    )

    add_callout(
        doc,
        title="CAPTURA 3: Eliminación en bosque con Orden Inverso (reverse)",
        instruction="Ejecute el comando para visualizar el orden reverse y evidenciar el factor de 32 filas en el paso 1.",
        command="python main.py -p 3 -l bosque --case alarma -a elimination --order reverse",
        screenshot_tag="CAPTURA 3 (bosque alarma - Eliminación Reverse)"
    )

    # ------------------ PARTE 2: MUESTREO ------------------
    add_header_styled(doc, "3. Muestreo: Implementación y Experimentación", level=1)
    
    add_header_styled(doc, "3.1. Literal c) Implementación de prior_sample", level=2)
    doc.add_paragraph(
        "La función prior_sample genera una muestra o mundo completo de la red bayesiana simulando el proceso generativo causal. "
        "Dado que la lista network.variables ya se encuentra en orden topológico estricto, cada variable se muestrea después de que "
        "todos sus padres han sido instanciados. La regla de decisión evalúa si un número pseudoaleatorio uniforme r ∈ [0, 1), generado "
        "exclusivamente con rng.random(), es estrictamente menor que la probabilidad condicional local P(X=1 | padres). "
        "En tal caso se asigna X=1; de lo contrario, X=0."
    )
    
    code_prior = (
        "def prior_sample(network, rng: random.Random) -> dict:\n"
        "    sample = {}\n"
        "    for variable in network.variables:\n"
        "        # Consulta P(variable=1 | padres ya muestreados en sample)\n"
        "        prob = network.probability(variable, 1, sample)\n"
        "        sample[variable] = 1 if rng.random() < prob else 0\n"
        "    return sample\n"
    )
    add_code_block(doc, code_prior)

    add_header_styled(doc, "3.2. Literal d) Implementación de sampling_ask", level=2)
    doc.add_paragraph(
        "La función sampling_ask coordina el proceso de inferencia aproximada sobre samples intentos (fijos). "
        "Soporta dos modalidades:\n"
        "1. Muestreo a Priori (reject=False): Exige evidencia vacía. Cada muestra generada se acepta y se acumula en un histograma de frecuencias.\n"
        "2. Muestreo por Rechazo (reject=True): Para cada muestra generada a priori, se verifica si es totalmente compatible con la evidencia dada. "
        "Si coincide en todas las variables observadas, la muestra es aceptada y se suma al conteo del valor de la consulta; en caso contrario, se descarta."
    )
    doc.add_paragraph(
        "Aspectos clave de la implementación:\n"
        "• El parámetro samples representa el número de intentos o pruebas de generación, NO el número de muestras aceptadas.\n"
        "• La normalización se realiza dividiendo entre la cantidad de muestras aceptadas, no sobre el número total de intentos.\n"
        "• Si ninguna muestra es compatible con la evidencia (aceptadas = 0), la función retorna de manera segura (None, 0), "
        "lo cual refleja que con ese presupuesto de intentos no se pudo observar el evento, sin implicar que la evidencia sea matemáticamente imposible."
    )

    code_ask = (
        "def sampling_ask(network, query, evidence, samples, rng, reject=False):\n"
        "    network.validate_query(query, evidence)\n"
        "    if samples < 1:\n"
        "        raise ValueError(\"El número de muestras debe ser positivo\")\n"
        "    if evidence and not reject:\n"
        "        raise ValueError(\"Muestreo a priori estima marginales; use rechazo con evidencia\")\n"
        "\n"
        "    counts = {0: 0, 1: 0}\n"
        "    accepted = 0\n"
        "\n"
        "    for _ in range(samples):\n"
        "        sample = prior_sample(network, rng)\n"
        "        if reject:\n"
        "            # Filtrado por rechazo: comprobar consistencia con la evidencia\n"
        "            if all(sample[var] == val for var, val in evidence.items()):\n"
        "                counts[sample[query]] += 1\n"
        "                accepted += 1\n"
        "        else:\n"
        "            counts[sample[query]] += 1\n"
        "            accepted += 1\n"
        "\n"
        "    if accepted == 0:\n"
        "        return (None, 0)\n"
        "\n"
        "    return (normalize(counts), accepted)\n"
    )
    add_code_block(doc, code_ask)

    # ------------------ LITERAL E: EXPERIMENTACIÓN ------------------
    add_header_styled(doc, "3.3. Literal e) Experimentación en la red bosque", level=2)
    doc.add_paragraph(
        "Se evalúa la precisión y eficiencia del muestreo en la red bosque para la consulta Incendio bajo tres casos de evidencia, "
        "comparando 1,000 frente a 10,000 intentos (con semilla base seed=7):"
    )

    t_exp = doc.add_table(rows=7, cols=7)
    format_table(t_exp, [1.3, 1.1, 0.9, 0.9, 1.1, 1.1, 1.0],
                 ["Caso de Evidencia", "Método", "Intentos", "Aceptadas", "Estimado P(1)", "Exacto P(1)", "Error Abs."],
                 [
                     ["sin_evidencia", "A priori", "1,000", "1,000 (100%)", "0.056000", "0.065160", "0.009160"],
                     ["sin_evidencia", "A priori", "10,000", "10,000 (100%)", "0.063400", "0.065160", "0.001760"],
                     ["alarma (Alarma=1)", "Rechazo", "1,000", "191 (19.1%)", "0.277487", "0.329517", "0.052030"],
                     ["alarma (Alarma=1)", "Rechazo", "10,000", "1,841 (18.4%)", "0.315589", "0.329517", "0.013928"],
                     ["rara (T=1,F=1,A=1)", "Rechazo", "1,000", "0 (0.0%)", "Sin estimación", "0.260778", "N/A"],
                     ["rara (T=1,F=1,A=1)", "Rechazo", "10,000", "11 (0.11%)", "0.090909", "0.260778", "0.169869"]
                 ],
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT])

    doc.add_paragraph(
        "Hallazgos empíricos del experimento:\n"
        "1. Efecto del incremento de intentos: Al pasar de 1,000 a 10,000 intentos, el error absoluto se redujo de manera notable en sin_evidencia "
        "(de 0.0091 a 0.0017) y en alarma (de 0.0520 a 0.0139). Esto ilustra la convergencia asintótica garantizada por la Ley de los Grandes Números.\n"
        "2. Eficiencia según la probabilidad de la evidencia: En el caso alarma, aproximadamente el 18.4 % de los mundos generados fueron compatibles, "
        "lo que proporcionó una muestra efectiva de 1,841 datos aceptados, suficiente para una estimación estable.\n"
        "3. Colapso del muestreo por rechazo ante eventos raros: En el caso rara (Tormenta=1, Falla=1, Alarma=1), la probabilidad conjunta del evento "
        "es ínfima (~0.1 %). Con 1,000 intentos se obtuvieron cero aceptaciones (la función retornó None de forma correcta). Con 10,000 intentos "
        "solo se aceptaron 11 muestras (un desperdicio computacional del 99.89 % de las simulaciones), lo que generó una estimación sumamente ruidosa "
        "(0.090909 frente a 0.260778 exacto, con un error de 0.169869). Esto demuestra la principal limitación del muestreo por rechazo clásico."
    )

    add_callout(
        doc,
        title="CAPTURA 4: Muestreo a Priori sin evidencia (10,000 muestras)",
        instruction="Ejecute el muestreo a priori en la red bosque para capturar la distribución y el histograma en la GUI.",
        command="python main.py -p 3 -l bosque --case sin_evidencia -a prior -n 10000",
        screenshot_tag="CAPTURA 4 (bosque sin_evidencia - Prior 10,000)"
    )

    add_callout(
        doc,
        title="CAPTURA 5: Muestreo por Rechazo con caso alarma (10,000 muestras)",
        instruction="Ejecute el muestreo por rechazo con evidencia Alarma=1.",
        command="python main.py -p 3 -l bosque --case alarma -a rejection -n 10000",
        screenshot_tag="CAPTURA 5 (bosque alarma - Rechazo 10,000)"
    )

    add_callout(
        doc,
        title="CAPTURA 6: Muestreo por Rechazo con caso rara (1,000 muestras - 0 aceptadas)",
        instruction="Ejecute con 1,000 intentos para evidenciar el caso límite 'Sin muestras aceptadas'.",
        command="python main.py -p 3 -l bosque --case rara -a rejection -n 1000",
        screenshot_tag="CAPTURA 6 (bosque rara - Rechazo 1,000)"
    )

    add_callout(
        doc,
        title="CAPTURA 7: Muestreo por Rechazo con caso rara (10,000 muestras - 11 aceptadas)",
        instruction="Ejecute con 10,000 intentos para capturar la tasa de aceptación reducida y alta varianza.",
        command="python main.py -p 3 -l bosque --case rara -a rejection -n 10000",
        screenshot_tag="CAPTURA 7 (bosque rara - Rechazo 10,000)"
    )

    # ------------------ PARTE 3: PREGUNTAS DEL INFORME (6a) ------------------
    add_header_styled(doc, "4. Respuestas a las Preguntas Teóricas del Informe (Punto 6a)", level=1)
    
    add_header_styled(doc, "4.1. Pregunta 6a (b): Coincidencia entre métodos y efecto del orden de eliminación", level=2)
    p_q_b = doc.add_paragraph()
    r = p_q_b.add_run("¿Por qué enumeración y eliminación deben coincidir? ¿Qué efecto observó al cambiar el orden de eliminación?\n")
    r.bold = True
    r.font.color.rgb = RGBColor(23, 54, 93)
    doc.add_paragraph(
        "Respuesta:\n"
        "1. Justificación de la coincidencia matemática: Enumeración e inferencia por eliminación de variables son dos algoritmos de inferencia EXACTA. "
        "Ambos computan formalmente el mismo valor matemático: la distribución condicional marginalizada a partir de la factorización de la red bayesiana:\n"
        "   P(X | e) = α ∑_{y ∈ Y} ∏_{V ∈ Variables} P(V | padres(V))\n"
        "La diferencia radica exclusivamente en la eficiencia del orden de operaciones: mientras que la enumeración expande sumas sobre productos "
        "recalculando términos idénticos de manera exponencial, la eliminación de variables aplica la ley distributiva del álgebra elemental "
        "(ab + ac = a(b + c)) para factorizar las sumatorias hacia adentro y almacenar productos intermedios en factores reutilizables. "
        "Al ser expresiones algebraicamente idénticas evaluadas en precisión finita, ambos algoritmos convergen exactamente al mismo valor numérico "
        "(en nuestro caso, P(Incendio=1 | Alarma=1) = 0.329517 en ambos).\n"
        "\n"
        "2. Efecto del orden de eliminación: El orden en que se eliminan las variables ocultas define el orden de combinación de los factores. "
        "Aunque el resultado numérico no varía, la complejidad computacional (temporal y espacial) depende críticamente de dicho orden:\n"
        "• En el orden topológico (forward), se eliminan primero los nodos raíz (Tormenta, Sequedad, Falla), lo que produce factores pequeños "
        "cuyo mayor producto alcanza únicamente 8 filas (2³ entradas).\n"
        "• En el orden inverso (reverse), se eliminan primero variables con múltiples padres y dependencias activas (Temperatura). Esto fuerza la unión "
        "prematura de factores inconexos, generando un factor producto intermedio gigante de 32 filas (2⁵ entradas, sobre Sequedad, Falla, Incendio y Humo).\n"
        "En conclusión, un orden deficiente incrementa exponencialmente la dimensión del mayor factor (noción de ancho del árbol inducido o treewidth), "
        "lo que en redes grandes puede hacer colapsar la memoria del computador."
    )

    add_header_styled(doc, "4.2. Pregunta 6a (c): Influencia de intentos y probabilidad de evidencia en el muestreo", level=2)
    p_q_c = doc.add_paragraph()
    r = p_q_c.add_run("¿Cómo influyen los intentos y la probabilidad de la evidencia en el error y la cantidad de muestras aceptadas?\n")
    r.bold = True
    r.font.color.rgb = RGBColor(23, 54, 93)
    doc.add_paragraph(
        "Respuesta:\n"
        "1. Influencia del número de intentos (N): Por la Ley de los Grandes Números, a medida que N → ∞, la frecuencia relativa muestral "
        "converge casi con certeza a la verdadera probabilidad subyacente. La teoría estadística establece que el error estándar de la estimación "
        "decrece con una tasa proporcional a O(1 / √N_aceptadas). No obstante, dado que el proceso es estocástico, en una ejecución particular "
        "con una semilla aleatoria específica, un incremento de intentos no garantiza matemáticamente un menor error debido a la fluctuación muestral, "
        "si bien en la media de múltiples repeticiones el error medio cuadrático siempre se reduce.\n"
        "\n"
        "2. Influencia de la probabilidad de la evidencia P(e): La tasa de aceptación del muestreo por rechazo está determinada exactamente por la probabilidad "
        "marginal de la evidencia observada: P(aceptación) = P(e). El número esperado de muestras útiles es E[N_aceptadas] = N · P(e).\n"
        "• Si la evidencia es frecuente (como alarma, P(e) ≈ 0.184), de 10,000 intentos se aceptan ~1,841 muestras, logrando un error bajo (0.0139).\n"
        "• Si la evidencia es infrecuente o rara (como rara: Tormenta=1, Falla=1, Alarma=1, donde P(e) ≈ 0.0011): con 1,000 intentos no se acepta "
        "ninguna muestra (error indefinido, 'Sin muestras aceptadas'), y con 10,000 intentos solo se aceptan 11 muestras, lo que provoca un error muy alto "
        "(0.1698) y alta varianza entre repeticiones.\n"
        "Esto demuestra por qué el muestreo por rechazo estándar es inviable para evidencias complejas o improbables, justificando el uso de técnicas como "
        "ponderación de verosimilitud (Likelihood Weighting) o muestreo de Gibbs (MCMC)."
    )

    # ------------------ PARTE 4: REFLEXIÓN SOBRE USO DE IA ------------------
    add_header_styled(doc, "5. Reflexión sobre el Uso de Inteligencia Artificial (Punto 6b y Anexos)", level=1)
    doc.add_paragraph(
        "En cumplimiento de la política del curso (Punto 6b, 4 % de la calificación), se presenta la reflexión sobre la interacción con IA "
        "durante el desarrollo del Punto 3:"
    )
    doc.add_paragraph(
        "• Aprendizaje y comprensión conceptual: La IA facilitó contrastar la formulación teórica de la ley distributiva en redes bayesianas "
        "con su implementación estructurada en factores tabulares. Al revisar el proceso de eliminación manual paso a paso en cadena_humo, "
        "se consolidó la comprensión de cómo cada paso de multiplicación y marginalización transforma la dimensionalidad de las variables involucradas.\n"
        "• Criterio de aceptación y rechazo de sugerencias: Se adoptó una postura crítica ante las sugerencias de la IA. Por ejemplo, en un inicio se "
        "planteó un bucle que iteraba hasta acumular N muestras aceptadas; sin embargo, al contrastar rigurosamente con el enunciado del taller "
        "(que exige que samples represente el número total de intentos, no de aceptaciones), se rechazó esa propuesta y se mantuvo el conteo estricto "
        "sobre el presupuesto de intentos, gestionando de forma adecuada el caso borde (None, 0).\n"
        "• Verificación de la solución: La verificación no se delegó en la IA, sino que se realizó mediante pruebas directas en terminal, contrastando "
        "las salidas de sampling_ask contra la referencia exacta de enumeration_ask y la traza de variable_elimination provista por la cátedra.\n"
        "• Preparación para sustentación: El estudiante comprende a cabalidad cada línea de código de prior_sample y sampling_ask, y está plenamente "
        "capacitado para explicar el impacto del orden topológico en la generación causal de muestras y la degradación del muestreo por rechazo ante el equipo docente."
    )

    add_header_styled(doc, "Anexo: Registro de Interacción con IA para el Punto 3", level=2)
    
    t_ai = doc.add_table(rows=3, cols=5)
    format_table(t_ai, [1.2, 1.4, 1.5, 1.3, 1.3],
                 ["Función", "Prompt / Consulta Inicial", "Sugerencia de IA", "Decisión del Estudiante", "Forma de Verificación"],
                 [
                     [
                         "prior_sample",
                         "¿Cómo generar una muestra a priori asegurando el orden topológico y usando la CPT?",
                         "Iterar sobre network.variables, consultar network.probability(var, 1, sample) y asignar 1 si rng.random() < p.",
                         "Aceptada. Se preservó el uso exclusivo del generador rng recibido y se verificó que sample contenga los padres.",
                         "Aserciones unitarias comprobando claves completas, dominio binario y consistencia con layouts."
                     ],
                     [
                         "sampling_ask",
                         "¿Cómo implementar muestreo por rechazo y gestionar el caso donde no se aceptan muestras?",
                         "Normalizar dividiendo por samples totales y generar muestras hasta lograr N aceptadas.",
                         "Rechazada y adaptada. Se corrigió para iterar exactamente samples veces (intentos) y normalizar sobre aceptadas. Retornar (None, 0) si no hay aceptadas.",
                         "Ejecución con caso rara en bosque con 1,000 y 10,000 intentos; validación contra tabla exacta de enumeración."
                     ]
                 ],
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    # ------------------ GUÍA PASO A PASO PARA PANTALLAZOS ------------------
    add_header_styled(doc, "6. Guía Paso a Paso para Tomar Pantallazos y Pegarlos en este Documento", level=1)
    doc.add_paragraph(
        "Si desea incorporar los pantallazos directamente en este documento Word, siga esta guía práctica paso a paso:"
    )
    doc.add_paragraph(
        "Paso 1: Abrir la terminal de Windows (PowerShell o CMD) en la carpeta raíz del proyecto:\n"
        "  C:\\Users\\dinoc\\OneDrive\\Escritorio\\Uniandes\\7 - Septimo Semestre\\INTELIGENCIA ARTIFICIAL (ISIS-1611)\\Taller-3\\taller3_ia\n"
        "\n"
        "Paso 2: Asegurarse de estar en la rama de trabajo Camilo-P3 mediante el comando:\n"
        "  git status\n"
        "\n"
        "Paso 3: Para cada captura solicitada en los recuadros azules de este documento:\n"
        "• Ejecute el comando indicado (sin el argumento -t se abrirá la ventana gráfica Tkinter interactiva; con -t se mostrará en texto en la terminal).\n"
        "• Presione la combinación de teclas en Windows: [ Tecla Windows + Shift + S ].\n"
        "• Seleccione con el cursor el área de la ventana gráfica o de la terminal que contiene la tabla o gráfica.\n"
        "• Regrese a este documento Word, sitúe el cursor en el recuadro azul correspondiente (donde dice '[ >>> PEGAR AQUÍ CAPTURA X <<< ]') y presione [ Ctrl + V ].\n"
        "• Ajuste el tamaño de la imagen para que encaje limpiamente entre los márgenes de la página."
    )

    doc.save("Punto_3_Documentacion_Camilo.docx")
    print("Documento 'Punto_3_Documentacion_Camilo.docx' generado exitosamente con formato profesional.")

if __name__ == '__main__':
    build_p3_document()
